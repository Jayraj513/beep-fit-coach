"""
Unit tests for document_parser.py
Verifies PDF, DOCX, empty, corrupted, and edge case extraction.
"""

import io
from document_parser import extract_document_text, format_file_size
import pypdf
import docx

def create_sample_pdf(text: str) -> bytes:
    from pypdf import PdfWriter
    from reportlab.pdfgen import canvas
    # Alternatively, create simple PDF using pypdf or reportlab, or if reportlab is not present, use raw minimal PDF
    pass

def main():
    print("Testing document_parser...")

    # 1. Test size formatting
    assert format_file_size(500) == "500 B"
    assert format_file_size(2048) == "2.0 KB"
    assert format_file_size(5 * 1024 * 1024) == "5.0 MB"
    print("[PASS] format_file_size passed")

    # 2. Test empty file
    text, err, meta = extract_document_text(b"", "empty.pdf")
    assert err and "empty" in err.lower()
    print("[PASS] Empty file test passed")

    # 3. Test unsupported extension
    text, err, meta = extract_document_text(b"some text", "test.txt")
    assert err and "unsupported" in err.lower()
    print("[PASS] Unsupported format test passed")

    # 4. Test legacy .doc
    text, err, meta = extract_document_text(b"some text", "test.doc")
    assert err and "legacy" in err.lower()
    print("[PASS] Legacy .doc warning passed")

    # 5. Test oversized file
    large_bytes = b"0" * (11 * 1024 * 1024)
    text, err, meta = extract_document_text(large_bytes, "big.pdf", max_size_mb=10)
    assert err and "exceeds" in err.lower()
    print("[PASS] Oversized file check passed")

    # 6. Test corrupted PDF
    corrupt_bytes = b"%PDF-1.4 garbage corrupt binary content \x00\xff"
    text, err, meta = extract_document_text(corrupt_bytes, "corrupt.pdf")
    assert err and ("unable to read" in err.lower() or "error" in err.lower())
    print("[PASS] Corrupted PDF handling passed")

    # 7. Test corrupted DOCX
    corrupt_docx = b"PK\x03\x04 fake zip header corrupted"
    text, err, meta = extract_document_text(corrupt_docx, "corrupt.docx")
    assert err and ("unable to read" in err.lower() or "not a word file" in err.lower())
    print("[PASS] Corrupted DOCX handling passed")

    # 8. Test valid DOCX creation and extraction
    doc = docx.Document()
    doc.add_heading("Anand Sharma - Resume", 0)
    doc.add_paragraph("Education: B.Tech Computer Science, 2024")
    doc.add_paragraph("Skills: Python, SQL, Git, Tableau")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Project"
    table.cell(0, 1).text = "Tech Stack"
    table.cell(1, 0).text = "Sales Dashboard"
    table.cell(1, 1).text = "Python, Streamlit"
    
    docx_buf = io.BytesIO()
    doc.save(docx_buf)
    docx_bytes = docx_buf.getvalue()

    text, err, meta = extract_document_text(docx_bytes, "anand_resume.docx")
    assert err is None, f"Expected no error, got: {err}"
    assert "Anand Sharma" in text
    assert "Python, SQL, Git" in text
    assert "Sales Dashboard" in text
    assert meta["char_count"] > 100
    print(f"[PASS] Valid DOCX extraction passed ({meta['char_count']} chars, {meta['num_paragraphs']} paragraphs, {meta['num_tables']} tables)")

    # 9. Test scanned-like PDF (PDF with 1 blank page without text)
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    pdf_buf = io.BytesIO()
    writer.write(pdf_buf)
    blank_pdf_bytes = pdf_buf.getvalue()

    text, err, meta = extract_document_text(blank_pdf_bytes, "scanned_invoice.pdf")
    assert err and ("scanned" in err.lower() or "ocr" in err.lower())
    assert meta["is_scanned"] is True
    print("[PASS] Scanned / image-only PDF detection passed")

    print("\nAll parser tests passed successfully!")

if __name__ == "__main__":
    main()
