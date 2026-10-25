"""
Document extraction module for Beep Fit Coach.
Supports PDF and DOCX files with validation, scanned document detection, and memory-safe processing.
"""

import io
import os
from typing import Tuple, Dict, Any

MAX_FILE_SIZE_MB_DEFAULT = 10


def format_file_size(size_bytes: int) -> str:
    """Formats bytes into human-readable size string."""
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    else:
        return f"{size_bytes / (1024 * 1024):.1f} MB"


def extract_text_from_pdf(stream: io.BytesIO) -> Tuple[str, str | None, Dict[str, Any]]:
    """
    Extracts text from a PDF stream using pypdf.
    Detects encrypted files, corrupted files, and image-only/scanned PDFs.
    """
    import pypdf
    import pypdf.errors

    metadata = {"num_pages": 0, "is_scanned": False}

    try:
        reader = pypdf.PdfReader(stream)
        num_pages = len(reader.pages)
        metadata["num_pages"] = num_pages

        if num_pages == 0:
            return "", "The PDF document contains 0 pages.", metadata

        # Check for password encryption
        if reader.is_encrypted:
            try:
                decrypted = reader.decrypt("")
                if decrypted == pypdf.constants.DecryptionResult.FAILED or reader.is_encrypted:
                    return (
                        "",
                        "This PDF file is password-protected or encrypted. Please remove the password protection or copy and paste the text directly.",
                        metadata,
                    )
            except Exception:
                return (
                    "",
                    "This PDF file is password-protected or encrypted. Please remove the password protection or copy and paste the text directly.",
                    metadata,
                )

        extracted_pages = []
        for idx, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    extracted_pages.append(page_text.strip())
            except Exception:
                continue

        full_text = "\n\n".join(extracted_pages).strip()

        # Scanned PDF detection:
        # If the PDF has pages but the total extracted text is under 30 characters
        if len(full_text) < 30 and num_pages > 0:
            metadata["is_scanned"] = True
            return (
                "",
                "**Notice: No selectable text found in this PDF.**\n\n"
                "This document appears to be a scanned image or photo-based PDF without embedded digital text. "
                "Optical Character Recognition (OCR) is required to read scanned documents. "
                "Please upload a standard text-based PDF or Word document, or copy and paste your text directly.",
                metadata,
            )

        return full_text, None, metadata

    except pypdf.errors.PdfReadError as err:
        return (
            "",
            f"Unable to read this PDF file. The document may be corrupted or improperly formatted. Details: {str(err)[:120]}",
            metadata,
        )
    except Exception as general_err:
        return (
            "",
            f"An error occurred while parsing the PDF: {str(general_err)[:120]}. Please try saving it again or copy/paste the text.",
            metadata,
        )


def extract_text_from_docx(stream: io.BytesIO) -> Tuple[str, str | None, Dict[str, Any]]:
    """
    Extracts text from a DOCX stream using python-docx.
    Extracts both paragraphs and table cells to preserve structured resume data.
    """
    import docx

    metadata = {"num_paragraphs": 0, "num_tables": 0}

    try:
        doc = docx.Document(stream)
        text_elements = []

        # Extract paragraphs
        metadata["num_paragraphs"] = len(doc.paragraphs)
        for p in doc.paragraphs:
            p_text = p.text.strip()
            if p_text:
                text_elements.append(p_text)

        # Extract table contents (many resumes use tables for layout)
        metadata["num_tables"] = len(doc.tables)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                # Deduplicate identical adjacent cells that result from merged table cells
                deduped = []
                for cell in row_cells:
                    if not deduped or cell != deduped[-1]:
                        deduped.append(cell)
                if deduped:
                    text_elements.append(" | ".join(deduped))

        full_text = "\n\n".join(text_elements).strip()

        if len(full_text) < 20:
            return (
                "",
                "The Word document contains very little or no readable text. Please check the document content or paste the text directly.",
                metadata,
            )

        return full_text, None, metadata

    except Exception as docx_err:
        err_str = str(docx_err)
        if "Package not found" in err_str or "not a Word file" in err_str.lower():
            return (
                "",
                "Unable to open this Word document. It may be a legacy binary '.doc' file (pre-2007) or corrupted. "
                "Please save it as a modern '.docx' file or copy and paste the text directly.",
                metadata,
            )
        return (
            "",
            f"Unable to read this Word document: {err_str[:120]}. Please save it again or copy/paste the text.",
            metadata,
        )


def extract_document_text(
    file_bytes: bytes,
    filename: str,
    max_size_mb: int = MAX_FILE_SIZE_MB_DEFAULT,
) -> Tuple[str, str | None, Dict[str, Any]]:
    """
    Universal document extraction entry point.
    Validates file size, file format, extracts text in-memory, and returns structured result.
    
    Returns:
        (extracted_text, error_message, metadata_dict)
    """
    file_size = len(file_bytes)
    formatted_size = format_file_size(file_size)
    base_metadata: Dict[str, Any] = {
        "filename": filename,
        "file_size_bytes": file_size,
        "formatted_size": formatted_size,
        "char_count": 0,
        "word_count": 0,
    }

    # 1. Check empty file
    if file_size == 0:
        return (
            "",
            "The uploaded file is empty (0 bytes). Please upload a valid document or paste text directly.",
            base_metadata,
        )

    # 2. Check maximum size
    max_bytes = max_size_mb * 1024 * 1024
    if file_size > max_bytes:
        return (
            "",
            f"The uploaded file ({formatted_size}) exceeds the maximum limit of {max_size_mb} MB. "
            f"Please compress the document or copy/paste the text directly.",
            base_metadata,
        )

    # 3. Detect format
    ext = os.path.splitext(filename)[1].lower()
    stream = io.BytesIO(file_bytes)

    if ext == ".pdf":
        text, error, doc_meta = extract_text_from_pdf(stream)
    elif ext == ".docx":
        text, error, doc_meta = extract_text_from_docx(stream)
    elif ext == ".doc":
        return (
            "",
            "Legacy '.doc' files (Word 97-2003) are not supported. Please convert the file to '.docx' or '.pdf', or paste the text directly.",
            base_metadata,
        )
    else:
        return (
            "",
            f"Unsupported file format '{ext}'. Only PDF (.pdf) and Word documents (.docx) are supported. You can also paste text directly.",
            base_metadata,
        )

    base_metadata.update(doc_meta)
    if text:
        base_metadata["char_count"] = len(text)
        base_metadata["word_count"] = len(text.split())

    return text, error, base_metadata
